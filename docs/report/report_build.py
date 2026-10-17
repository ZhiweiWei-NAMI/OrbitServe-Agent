#!/usr/bin/env python3
"""Rebuild the reviewed Chinese report from evidence-backed report_model.json.
Dependencies: python-docx. Images remain relative to this source directory.
"""
from pathlib import Path
import argparse,json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
BASE=Path(__file__).resolve().parent

def font(run,size=None,bold=None):
    run.font.name='Times New Roman'
    rf=run._element.get_or_add_rPr().get_or_add_rFonts()
    for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('cs','Times New Roman'),('eastAsia','宋体')]: rf.set(qn('w:'+k),v)
    for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']: rf.attrib.pop(qn('w:'+k),None)
    run.font.color.rgb=RGBColor(0,0,0)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold

def para(doc,text='',style=None,bold=None,align=None):
    p=doc.add_paragraph(style=style)
    if text: font(p.add_run(text),bold=bold)
    if align is not None: p.alignment=align
    return p

def no_split(row):
    prop=row._tr.get_or_add_trPr(); el=OxmlElement('w:cantSplit'); prop.append(el)

def table(doc,headers,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    if widths:
        for col,w in zip(t.columns,widths): col.width=Cm(w)
    for i,h in enumerate(headers): t.rows[0].cells[i].text=str(h)
    trPr=t.rows[0]._tr.get_or_add_trPr();r=OxmlElement('w:tblHeader');r.set(qn('w:val'),'true');trPr.append(r)
    for rowdata in rows:
        cs=t.add_row().cells
        for i,v in enumerate(rowdata): cs[i].text=str(v)
    for ri,row in enumerate(t.rows):
        no_split(row)
        for ci,c in enumerate(row.cells):
            if widths:c.width=Cm(widths[ci])
            c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr=c._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
            for k in ['top','bottom','left','right']:
                e=OxmlElement('w:'+k);e.set(qn('w:w'),'70' if k in ['top','bottom'] else '90');e.set(qn('w:type'),'dxa');margin.append(e)
            tcPr.append(margin)
            borders=OxmlElement('w:tcBorders')
            for k in ['top','bottom','left','right']:
                e=OxmlElement('w:'+k);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D0D0D0');borders.append(e)
            tcPr.append(borders)
            if ri==0:
                shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F2F2F2');tcPr.append(shd)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0)
                p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
                p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.EXACTLY;p.paragraph_format.line_spacing=Pt(22)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                if len(rows)<=4:p.paragraph_format.keep_with_next=not(ri==len(rows) and ci==len(headers)-1)
                for run in p.runs:font(run,10.5,bold=(ri==0))
    return t

def build(model,output):
    doc=Document();sec=doc.sections[0]
    sec.page_height=Cm(29.7);sec.page_width=Cm(21)
    sec.top_margin=Cm(2.4);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(2.6);sec.right_margin=Cm(2.6)
    sec.header_distance=Cm(1.1);sec.footer_distance=Cm(1.1)
    for style in doc.styles:
        for border in list(style.element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
    sizes={'Normal':12,'Title':21,'Subtitle':14,'Heading 1':16,'Heading 2':14,'Heading 3':12,'Caption':10.5}
    for name,size in sizes.items():
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0);s.font.italic=False;s.font.underline=False
        rf=s.element.get_or_add_rPr().get_or_add_rFonts()
        for k,v in [('ascii','Times New Roman'),('hAnsi','Times New Roman'),('cs','Times New Roman'),('eastAsia','宋体')]:rf.set(qn('w:'+k),v)
        for k in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:rf.attrib.pop(qn('w:'+k),None)
        p=s.paragraph_format;p.line_spacing_rule=WD_LINE_SPACING.EXACTLY;p.line_spacing=Pt(22)
        p.space_after=Pt(0);p.space_before=Pt(0)
        p.widow_control=True
        if name=='Normal':p.first_line_indent=Cm(.84)
        else:p.first_line_indent=Cm(0)
        if name.startswith('Heading'):p.space_before=Pt(10);p.space_after=Pt(6);p.keep_with_next=True;s.font.bold=True
        if name=='Title':p.space_before=Pt(18);p.space_after=Pt(15);p.line_spacing=Pt(33);p.keep_with_next=True
        if name=='Subtitle':p.space_after=Pt(12)
        if name=='Caption':p.space_before=Pt(4);p.space_after=Pt(5);p.keep_with_next=False
    doc.core_properties.title=model['title'];doc.core_properties.subject='遥感任务图自适应分解与协同执行的研究总结及可复现实验'
    doc.core_properties.author='OrbitServe-Agent';doc.core_properties.keywords='遥感;任务语义;可执行图;预测调度;协同卸载;可复现'
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for v in ['第 ',None,' 页']:
        run=footer.add_run(v or '');font(run,10)
        if v is None:
            field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');run._r.addnext(field)
    records=[];pending_page_break=False
    for item in model['content']:
        typ=item['type'];records.append(item)
        if typ=='pagebreak':pending_page_break=True
        elif typ=='title':para(doc,item['text'],'Title',align=WD_ALIGN_PARAGRAPH.CENTER)
        elif typ=='subtitle':para(doc,item['text'],'Subtitle',align=WD_ALIGN_PARAGRAPH.CENTER)
        elif typ=='heading':
            hp=para(doc,item['text'],'Heading '+str(item.get('level',1)))
            if pending_page_break:hp.paragraph_format.page_break_before=True;pending_page_break=False
        elif typ=='p':para(doc,item['text'])
        elif typ=='note':
            p=para(doc,item['text']);p.paragraph_format.first_line_indent=Cm(0)
            for run in p.runs:font(run,10.5)
        elif typ=='equation':
            p=para(doc,item['text'],align=WD_ALIGN_PARAGRAPH.CENTER);p.paragraph_format.first_line_indent=Cm(0)
        elif typ=='table':
            cap=para(doc,item['caption'],'Caption',align=WD_ALIGN_PARAGRAPH.CENTER);cap.paragraph_format.keep_with_next=True
            table(doc,item['headers'],item['rows'],item.get('widths'))
            if item.get('note'):
                p=para(doc,item['note']);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_together=True
                for run in p.runs:font(run,10.5)
        elif typ=='figure':
            path=BASE/item['path'];assert path.is_file(),path
            p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing_rule=WD_LINE_SPACING.SINGLE;p.paragraph_format.line_spacing=1
            p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(0);p.paragraph_format.keep_with_next=True
            shape=p.add_run().add_picture(str(path),width=Cm(item.get('width_cm',15.6)))
            shape._inline.docPr.set('descr',item.get('alt',item['caption']))
            cp=para(doc,item['caption'],'Caption',align=WD_ALIGN_PARAGRAPH.CENTER)
            cp.paragraph_format.keep_with_next=bool(item.get('note'))
            if item.get('note'):
                p=para(doc,item['note']);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.keep_together=True
                for run in p.runs:font(run,10.5)
        else:raise ValueError(typ)
    doc.save(output)
    # Human-editable textual source, consistent with the assembled report model.
    md=[]
    for x in records:
        t=x['type']
        if t in ['title','subtitle','heading']:md.append('#'*({'title':1,'subtitle':2}.get(t,x.get('level',1)+1))+' '+x['text'])
        elif t in ['p','note','equation']:md.append(x['text'])
        elif t=='figure':md.extend([f"![{x['caption']}]({x['path']})",(x.get('note') or '')])
        elif t=='table':
            md.append(x['caption']);md.append('| '+' | '.join(x['headers'])+' |');md.append('| '+' | '.join(['---']*len(x['headers']))+' |')
            md.extend('| '+' | '.join(map(str,r))+' |' for r in x['rows']);md.append(x.get('note') or '')
    (BASE/'report_source.md').write_text('\n\n'.join(md),encoding='utf-8')
    print(output)
if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--model',default=str(BASE/'report_model.json'));ap.add_argument('--output',default=str(BASE/'OrbitServe_Agent_研究总结与实验报告.docx'));a=ap.parse_args()
    build(json.loads(Path(a.model).read_text()),Path(a.output))
